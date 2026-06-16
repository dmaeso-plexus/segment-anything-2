#!/usr/bin/env python3
"""Generate the Nexus42 AI 'AI Tech Lead' offer letter as an editable .docx.

Same fictional content as build_offer.py (fictional Nexus42 company, original
logo, fictional signatory) — just in Word format so the text can be edited.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "nexus42-logo.png")
SIGN = "/tmp/sign_seal.png"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
GOLD = RGBColor(0xC9, 0xA2, 0x4B)
INK = RGBColor(0x1B, 0x27, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def make_sign_seal():
    """Render the cursive signature + company seal to a transparent PNG."""
    from reportlab.pdfgen import canvas
    import build_offer as bo
    import fitz
    from PIL import Image
    c = canvas.Canvas("/tmp/sign_seal.pdf", pagesize=(320, 120))
    bo.draw_signature(c, 12, 64, 1.0)
    bo.draw_stamp(c, 250, 58, 32, angle=-9)
    c.save()
    pix = fitz.open("/tmp/sign_seal.pdf")[0].get_pixmap(dpi=600, alpha=True)
    pix.save(SIGN)
    im = Image.open(SIGN).convert("RGBA")
    im = im.crop(im.getbbox())
    pad = int(im.height * 0.06)
    cv = Image.new("RGBA", (im.width + 2 * pad, im.height + 2 * pad), (0, 0, 0, 0))
    cv.paste(im, (pad, pad), im)
    cv.save(SIGN)


def shade(cell, hexstr):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexstr)
    tcPr.append(shd)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def bottom_rule(paragraph, color="0B1F3A", sz="14"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), color)
    pbdr.append(bottom); pPr.append(pbdr)


def page_field(paragraph):
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"
    r.append(t); fld.append(r); paragraph._p.append(fld)


def para(doc, text, size=10, color=INK, justify=True, after=7, before=0, bold_all=False):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(size); r.font.color.rgb = color
        r.bold = bold_all or (i % 2 == 1)
    return p


def heading(doc, text):
    return para(doc, "**" + text + "**", size=11, color=NAVY, justify=False,
                after=2, before=8)


def cellset(cell, text, bold=False, right=False, white=False, size=9, color=INK):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    r.font.color.rgb = WHITE if white else color


def build(path):
    make_sign_seal()
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(1.9))
    CW = Cm(17.2)  # content width

    # ---- letterhead (logo + company block) ----
    lt = doc.add_table(rows=1, cols=2); no_borders(lt); lt.autofit = False
    lt.columns[0].width = Cm(8.6); lt.columns[1].width = Cm(8.6)
    lc, rc = lt.rows[0].cells
    lc.width = Cm(8.6); rc.width = Cm(8.6)
    lc.paragraphs[0].add_run().add_picture(LOGO, width=Inches(2.15))
    rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = rp.add_run("Nexus42 Artificial Intelligence Holding PJSC")
    run.bold = True; run.font.size = Pt(9); run.font.color.rgb = NAVY
    for line in ("Capital Gate, 12th Floor, Al Khaleej Al Arabi Street",
                 "Abu Dhabi, United Arab Emirates",
                 "people@nexus42.ai  ·  +971 2 555 0042  ·  nexus42.ai"):
        rp.add_run("\n" + line).font.size = Pt(7.5)
        rp.runs[-1].font.color.rgb = GREY
    rule = doc.add_paragraph(); rule.paragraph_format.space_after = Pt(8)
    bottom_rule(rule)

    # ---- ref / date / confidential ----
    para(doc, "**PRIVATE & CONFIDENTIAL**", size=9, color=GOLD, justify=False, after=2)
    rd = doc.add_paragraph(); rd.paragraph_format.space_after = Pt(8)
    rd.paragraph_format.tab_stops.add_tab_stop(CW, WD_TAB_ALIGNMENT.RIGHT)
    r1 = rd.add_run("Ref:  NX42/HR/OFR/2026/0418\tDate:  16 June 2026")
    r1.font.size = Pt(9); r1.font.color.rgb = GREY

    para(doc, "Diego Martín Maeso", size=9, color=GREY, justify=False, after=0)
    para(doc, "Avenida de Asturias 31, 1A", size=9, color=GREY, justify=False, after=0)
    para(doc, "34880 Guardo (Palencia)", size=9, color=GREY, justify=False, after=0)
    para(doc, "España", size=9, color=GREY, justify=False, after=10)

    para(doc, "Dear Diego Martín Maeso,", after=4)
    para(doc, "**RE: OFFER OF EMPLOYMENT — AI TECH LEAD**", size=13, color=NAVY,
         justify=False, after=8)

    para(doc, "We are delighted to extend to you an offer of employment with "
              "**Nexus42 Artificial Intelligence Holding PJSC** (“Nexus42” or the "
              "“Company”). Following our selection process, the Executive Committee "
              "has approved your appointment to a senior technical leadership role "
              "within our Foundation Models & AI Platform division. This letter sets "
              "out the principal terms of our offer.")

    heading(doc, "1.  Position and Reporting")
    para(doc, "You will be employed as **AI Tech Lead**, reporting to the Vice "
              "President, AI Engineering. You will lead a multidisciplinary team "
              "responsible for the design, training and large-scale deployment of "
              "the Company’s foundation models and applied AI platforms, and will "
              "set technical direction, architecture and engineering standards "
              "across the division.")

    heading(doc, "2.  Nature and Term of Contract")
    para(doc, "Your employment is offered on a fixed-term basis for an **initial "
              "period of eighteen (18) months** (the “Initial Term”), commencing on "
              "the Commencement Date and renewable by mutual written agreement. The "
              "contract is governed by UAE Federal Decree-Law No. 33 of 2021 (the "
              "UAE Labour Law) and its implementing regulations.")

    heading(doc, "3.  Commencement Date and Place of Work")
    para(doc, "Your Commencement Date is anticipated to be **1 September 2026**, or "
              "such other date as may be mutually agreed in writing. Your principal "
              "place of work will be the Company’s headquarters at Capital Gate, "
              "Al Khaleej Al Arabi Street, Abu Dhabi, United Arab Emirates, with "
              "hybrid working available in accordance with Company policy.")

    heading(doc, "4.  Probationary Period")
    para(doc, "The first six (6) months of the Initial Term shall constitute a "
              "probationary period, during which either party may terminate the "
              "employment in accordance with the UAE Labour Law.")

    heading(doc, "5.  Remuneration")
    para(doc, "Your **annual base salary** will be **AED 3,000,000** (three million "
              "United Arab Emirates Dirhams), equivalent to approximately "
              "**EUR 750,000** per annum at an indicative reference rate of "
              "EUR 1.00 = AED 4.00. **In addition to your base salary, and "
              "separately from it**, you will receive the cash allowances set out "
              "below. All amounts are paid in UAE Dirhams (AED), free of UAE "
              "personal income tax.", after=4)

    rows = [
        ("Component", "Annual (AED)", "Monthly (AED)"),
        ("Base salary", "3,000,000", "250,000"),
        ("Housing allowance (in addition to salary)", "600,000", "50,000"),
        ("Transport allowance (in addition to salary)", "120,000", "10,000"),
        ("Total annual cash package", "3,720,000", "310,000"),
    ]
    tbl = doc.add_table(rows=5, cols=3); tbl.style = "Table Grid"; tbl.autofit = False
    widths = (Cm(9.2), Cm(4.0), Cm(4.0))
    for i, (a, b, cc) in enumerate(rows):
        cells = tbl.rows[i].cells
        for j, w in enumerate(widths):
            cells[j].width = w
        if i == 0:
            for j, txt in enumerate((a, b, cc)):
                cellset(cells[j], txt, bold=True, white=True, right=(j > 0))
                shade(cells[j], "0B1F3A")
        else:
            tot = (i == 4)
            cellset(cells[0], a, bold=tot)
            cellset(cells[1], b, bold=tot, right=True)
            cellset(cells[2], cc, bold=tot, right=True)
            if tot:
                for j in range(3):
                    shade(cells[j], "EFF3F7")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    para(doc, "Your base salary and allowances are payable monthly in arrears by "
              "bank transfer to a UAE account nominated by you. The total annual "
              "cash package above (AED 3,720,000, approximately EUR 930,000) is "
              "exclusive of the benefits set out in clause 7, which are provided in "
              "addition.", size=9, color=GREY, before=4)

    heading(doc, "6.  Performance Bonus and Long-Term Incentive")
    para(doc, "You will be eligible for a discretionary annual performance bonus of "
              "up to 30% of your base salary, subject to Company and individual "
              "performance, and will be invited to participate in the Company’s "
              "Long-Term Incentive Plan, details of which will be provided "
              "separately.")

    heading(doc, "7.  Benefits")
    for b in ("30 calendar days of paid annual leave per year, in addition to UAE public holidays;",
              "comprehensive private medical insurance for you and your eligible dependents;",
              "annual return air tickets to your home country for you and your eligible dependents;",
              "a one-time relocation allowance and assistance with relocation to Abu Dhabi;",
              "UAE residence visa sponsorship for you and your eligible dependents;",
              "end-of-service gratuity calculated in accordance with the UAE Labour Law."):
        bp = doc.add_paragraph(b, style="List Bullet")
        bp.paragraph_format.space_after = Pt(2)
        for r in bp.runs:
            r.font.size = Pt(10); r.font.color.rgb = INK

    heading(doc, "8.  Working Hours")
    para(doc, "Standard working hours are 40 hours per week, Monday to Friday, "
              "exclusive of breaks, together with the flexibility expected of a "
              "senior leadership position.")

    heading(doc, "9.  Confidentiality, Intellectual Property and Restrictive Covenants")
    para(doc, "Your employment is conditional upon execution of the Company’s "
              "standard Confidentiality and Intellectual Property Assignment "
              "Agreement. All intellectual property created in the course of your "
              "employment shall vest in the Company. Reasonable post-termination "
              "confidentiality and non-solicitation obligations will apply.")

    heading(doc, "10.  Conditions Precedent")
    para(doc, "This offer is conditional upon: (a) satisfactory completion of "
              "background, reference and security checks; (b) evidence of your "
              "right to be sponsored and to work in the UAE; (c) a satisfactory "
              "medical fitness examination; and (d) provision of attested copies of "
              "your academic and professional certificates.")

    heading(doc, "11.  Validity and Acceptance")
    para(doc, "This offer is open for acceptance until **30 June 2026**. To accept, "
              "please sign and return the duplicate of this letter. This letter "
              "constitutes the principal terms of our offer and supersedes any "
              "prior discussions; a full employment contract will be issued upon "
              "acceptance.")
    para(doc, "We are confident that your expertise will make a significant "
              "contribution to Nexus42’s mission to build sovereign, world-class "
              "artificial intelligence, and we look forward to welcoming you to the "
              "team.", after=10)

    para(doc, "Yours sincerely,", after=2)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
    sp.add_run().add_picture(SIGN, width=Inches(2.6))
    para(doc, "**Khalid Al Nuaimi**", justify=False, after=0)
    para(doc, "Group Chief Executive Officer", size=9, color=GREY, justify=False, after=0)
    para(doc, "For and on behalf of Nexus42 Artificial Intelligence Holding PJSC",
         size=9, color=GREY, justify=False, after=12)

    acc = doc.add_paragraph(); acc.paragraph_format.space_after = Pt(6); bottom_rule(acc, "D8DEE6", "6")
    heading(doc, "ACCEPTANCE")
    para(doc, "I, **Diego Martín Maeso**, accept the offer of employment on the "
              "terms set out in this letter.", after=14)
    al = doc.add_paragraph()
    al.paragraph_format.tab_stops.add_tab_stop(CW, WD_TAB_ALIGNMENT.RIGHT)
    ar = al.add_run("Signature: ______________________________\tDate: __________________")
    ar.font.size = Pt(10); ar.font.color.rgb = INK

    # ---- footer ----
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(CW, WD_TAB_ALIGNMENT.RIGHT)
    fr = fp.add_run("Nexus42 Artificial Intelligence Holding PJSC  ·  Licence "
                    "CN-2310045  ·  Capital Gate, Al Khaleej Al Arabi Street, "
                    "Abu Dhabi, UAE\tPage ")
    fr.font.size = Pt(7); fr.font.color.rgb = GREY
    page_field(fp)
    f2 = sec.footer.add_paragraph("This document is private and confidential and "
                                  "is intended solely for the named addressee.")
    f2.runs[0].font.size = Pt(7); f2.runs[0].font.color.rgb = GREY

    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    build(os.path.join(HERE, "Nexus42_Offer_AI-Tech-Lead.docx"))
